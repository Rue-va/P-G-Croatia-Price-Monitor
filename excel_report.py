#!/usr/bin/env python3
"""
Builds docs/PG_Price_Report.xlsx — the downloadable version of the dashboard.
Run AFTER scrape_and_export.py (it reads the files that script writes), in
the same daily job, so the workbook and the dashboard describe the same day.

Sheets (each is a filterable Excel table with a frozen header row):
  Overview            one row per retailer: data date, stores, P&G range,
                      TDP, promotions, missing products, price position
  Price comparison    every P&G product (barcode) sold at 2+ retailers,
                      with each retailer's regular price and the cheapest
  Retailer price gap  how much more each retailer typically charges than
                      the cheapest retailer for the same P&G products
  Categories          per category and retailer: P&G share of products,
                      share of TDP, typical price per unit vs. competitors
  Brands              every tracked brand per category and retailer:
                      products, stores, TDP, price per unit, promo share
  Promotions          every product on a published promotion today
  Missing products    P&G products gone from a store's list 3+ days
  Store portfolio     per store: P&G products listed, core range present
  <Retailer> SKUs     every P&G + tracked competitor product at that retailer
  <Retailer> by store every P&G listing, store by store

The definitions match the dashboard's "How the numbers are calculated".
"""
from __future__ import annotations

import json
from collections import defaultdict
from statistics import median

from openpyxl import Workbook
from openpyxl.styles import Alignment, Font, PatternFill
from openpyxl.utils import get_column_letter
from openpyxl.worksheet.table import Table, TableStyleInfo

from config import DOCS_DIR

REPORT_PATH = DOCS_DIR / "PG_Price_Report.xlsx"

HEADER_FILL = PatternFill(start_color="004B93", end_color="004B93", fill_type="solid")
HEADER_FONT = Font(color="FFFFFF", bold=True)
TITLE_FONT = Font(bold=True, size=14, color="004B93")
NOTE_FONT = Font(italic=True, color="666666")
EUR = "€#,##0.00"
EUR3 = "€#,##0.000"
PCT = "0.0%"
INT = "#,##0"

UNIT_LABEL = {"l": "litre", "kg": "kg", "kom": "piece", "pranje": "wash"}


def retailer_name(retailer: str) -> str:
    return retailer.title()


def _load(path, default):
    try:
        return json.loads(path.read_text(encoding="utf-8"))
    except (FileNotFoundError, json.JSONDecodeError):
        return default


def _real_ean(ean: str | None) -> bool:
    return bool(ean) and ":" not in ean


def _is_national(r: dict) -> bool:
    """Dm publishes one national list instead of one per store."""
    return (r.get("stores_count") or 0) <= 1


_used_table_names: set[str] = set()


def _sheet_name(name: str) -> str:
    for ch in "[]:*?/\\":
        name = name.replace(ch, "")
    return name[:31]


def _table_name(title: str) -> str:
    base = "".join(ch if ch.isalnum() else "_" for ch in title)[:40] or "T"
    if not base[0].isalpha():
        base = "T_" + base
    name, i = base, 2
    while name in _used_table_names:
        name, i = f"{base}_{i}", i + 1
    _used_table_names.add(name)
    return name


def _add_table(wb, title: str, headers: list[str], rows: list[list], formats: dict[str, str] | None = None,
               widths: dict[str, int] | None = None, note: str | None = None):
    """One sheet = one Excel table: optional note line, styled header, frozen, filterable."""
    ws = wb.create_sheet(_sheet_name(title))
    start = 1
    if note:
        ws.append([note])
        ws["A1"].font = NOTE_FONT
        start = 2
    ws.append(headers)
    for cell in ws[start]:
        cell.fill, cell.font = HEADER_FILL, HEADER_FONT
        cell.alignment = Alignment(wrap_text=True, vertical="center")
    for row in rows:
        ws.append(row)
    ws.freeze_panes = ws.cell(row=start + 1, column=1)
    last_row = start + len(rows)
    if rows:
        table = Table(displayName=_table_name(title), ref=f"A{start}:{get_column_letter(len(headers))}{last_row}")
        table.tableStyleInfo = TableStyleInfo(name="TableStyleMedium2", showRowStripes=True)
        ws.add_table(table)
    formats, widths = formats or {}, widths or {}
    for idx, h in enumerate(headers, start=1):
        ws.column_dimensions[get_column_letter(idx)].width = widths.get(h, max(11, min(40, len(h) + 3)))
        fmt = formats.get(h)
        if fmt and rows:
            for (cell,) in ws.iter_rows(min_row=start + 1, max_row=last_row, min_col=idx, max_col=idx):
                cell.number_format = fmt
    return ws


def _norm(s: str | None) -> str:
    return " ".join((s or "").replace("-", " ").replace("&", " & ").split()).upper()


def _brand_match(sku: dict, brand: str, side: str) -> bool:
    if (side == "pg") != (sku["side"] == "pg"):
        return False
    b, sb = _norm(brand), _norm(sku.get("brand"))
    return sb == b or sb.startswith(b + " ") or b in _norm(sku.get("product"))


def build_report():
    data = _load(DOCS_DIR / "data.json", None)
    if not data:
        raise SystemExit("docs/data.json not found — run scrape_and_export.py first")
    retailers = [r for r in data.get("retailers", []) if r.get("status") != "error"]
    rids = [r["retailer"] for r in retailers]
    skus = {rid: _load(DOCS_DIR / "skus" / f"{rid}.json", []) for rid in rids}
    flags = {rid: _load(DOCS_DIR / "flags" / f"{rid}.json", []) for rid in rids}

    wb = Workbook()
    wb.remove(wb.active)
    _used_table_names.clear()

    # ---------- same barcode across retailers ----------
    by_ean: dict[str, dict] = {}
    for rid in rids:
        for s in skus[rid]:
            if s["side"] != "pg" or not _real_ean(s["ean"]) or not s.get("shelf"):  # skip missing / €0.00
                continue
            g = by_ean.setdefault(s["ean"], {"product": s["product"], "brand": s["brand"], "cat": s.get("cat"), "prices": {}})
            g["cat"] = g["cat"] or s.get("cat")
            g["prices"][rid] = s["shelf"]
    shared = []
    for ean, g in by_ean.items():
        if len(g["prices"]) < 2:
            continue
        lo, hi = min(g["prices"].values()), max(g["prices"].values())
        # >2.5x for the identical product = clearance/placeholder price in a file
        shared.append({**g, "ean": ean, "lo": lo, "hi": hi, "suspect": hi / lo > 2.5,
                       "cheapest": [rid for rid, v in g["prices"].items() if v == lo]})
    shared.sort(key=lambda g: (g["suspect"], -(g["hi"] / g["lo"])))

    ok_shared = [g for g in shared if not g["suspect"]]
    gap_rows, gap_by_rid = [], {}
    for rid in rids:
        extras = [g["prices"][rid] / g["lo"] - 1 for g in ok_shared if rid in g["prices"]]
        if not extras:
            continue
        cheapest = sum(1 for g in ok_shared if rid in g["cheapest"])
        gap_by_rid[rid] = median(extras)
        gap_rows.append([retailer_name(rid), len(extras), cheapest, cheapest / len(extras), median(extras)])
    gap_rows.sort(key=lambda r: r[4])

    # ---------- Overview ----------
    ws = wb.create_sheet("Overview")
    ws["A1"] = "P&G Croatia — Price & Assortment Monitor"
    ws["A1"].font = TITLE_FONT
    ws["A2"] = f"Price lists of {data.get('as_of_date')} · generated {data.get('generated_at', '')[:16].replace('T', ' ')} UTC"
    ws["A2"].font = NOTE_FONT
    headers = ["Retailer", "Data date", "Status", "Stores reporting", "P&G products listed", "P&G TDP (products × stores)",
               "P&G products on promotion", "P&G products missing (3+ days)", "Typically above cheapest retailer"]
    ws.append([])
    ws.append(headers)
    for cell in ws[4]:
        cell.fill, cell.font = HEADER_FILL, HEADER_FONT
        cell.alignment = Alignment(wrap_text=True, vertical="center")
    for r in retailers:
        rid = r["retailer"]
        pg = [s for s in skus[rid] if s["side"] == "pg"]
        ins = r.get("insights") or {}
        ws.append([
            retailer_name(rid), r.get("last_updated"),
            "today's list" if r.get("status") == "ok" else "last available list",
            "one national list" if _is_national(r) else r.get("stores_count"),
            len(pg),
            None if _is_national(r) else sum(s["stores"] for s in pg),
            ins.get("promos_pg_total"),
            None if _is_national(r) else ins.get("flagged_skus_total"),
            gap_by_rid.get(rid),
        ])
    end = 4 + len(retailers)
    for row in ws.iter_rows(min_row=5, max_row=end):
        row[5].number_format = INT
        row[8].number_format = PCT
    for i, w in enumerate([14, 12, 18, 16, 14, 18, 16, 18, 18], start=1):
        ws.column_dimensions[get_column_letter(i)].width = w
    ws.row_dimensions[4].height = 45
    notes = [
        "How to read this workbook",
        "• Source: retailers' legally required public daily price lists (NN 75/2025, NN 101/2026). Konzum, Kaufland, Spar, Lidl and Plodine publish one list per store; Dm publishes one national list.",
        "• Prices are the typical (median) regular shelf price across a retailer's stores; promotions are reported separately.",
        "• TDP (total distribution points) = for each product, the number of stores listing it, added up. Not applicable to Dm.",
        "• Missing products: listed by a store before, absent from its list 3+ days in a row on days that store published. A signal to check, not a confirmed out-of-stock.",
        "• Core product (Store portfolio): listed in at least half of that retailer's stores today — the dashboard's own rule, not an official must-stock list.",
        "• Comparisons only include the brands tracked per category (config.py); store brands are excluded.",
    ]
    for i, n in enumerate(notes):
        c = ws.cell(row=end + 3 + i, column=1, value=n)
        c.font = Font(bold=True) if i == 0 else NOTE_FONT

    # ---------- Price comparison / gap ----------
    headers = ["Product", "Brand", "Category", "EAN"] + [retailer_name(r) for r in rids] + \
              ["Cheapest retailer", "Cheapest price", "Highest price", "Difference", "Check"]
    rows = [[g["product"], g["brand"], g["cat"] or "", g["ean"]] + [g["prices"].get(r) for r in rids] +
            [", ".join(retailer_name(r) for r in g["cheapest"]), g["lo"], g["hi"], g["hi"] / g["lo"] - 1,
             "check — likely clearance/placeholder price" if g["suspect"] else ""] for g in shared]
    fm = {retailer_name(r): EUR for r in rids}
    fm.update({"Cheapest price": EUR, "Highest price": EUR, "Difference": PCT})
    _add_table(wb, "Price comparison", headers, rows, fm, {"Product": 42, "EAN": 15, "Check": 32},
               note="Same P&G product (same barcode) at each retailer — typical regular shelf price across its stores.")
    _add_table(wb, "Retailer price gap", ["Retailer", "Shared P&G products", "Cheapest on", "Cheapest share",
                                          "Typically above cheapest"], gap_rows,
               {"Cheapest share": PCT, "Typically above cheapest": PCT},
               note="For P&G products sold at 2+ retailers: how much more each retailer typically charges than the cheapest one (gaps above 2.5× left out).")

    # ---------- Categories / Brands ----------
    cat_rows, brand_rows = [], []
    for r in retailers:
        rid, national, n_stores = r["retailer"], _is_national(r), r.get("stores_count") or 0
        for cat, w in (r.get("category_brand_watchlist") or {}).items():
            brands = [(b, "pg") for b in w.get("pg_brands", [])] + [(b, "comp") for b in w.get("competitor_brands", [])]
            priced = [s for s in skus[rid] if s.get("cat") == cat and s.get("unit_price") is not None
                      and s.get("unit") and s["unit"] != "pak"]
            units = defaultdict(int)
            for s in priced:
                units[s["unit"]] += 1
            unit = max(units, key=units.get) if units else None
            in_unit = [s for s in priced if s["unit"] == unit]
            if in_unit:  # same outlier rule as the dashboard's shelf view
                cm = median(s["unit_price"] for s in in_unit)
                in_unit = [s for s in in_unit if cm / 6 < s["unit_price"] < cm * 6]
            pg_vals = [s["unit_price"] for s in in_unit if s["side"] == "pg"]
            comp_vals = [s["unit_price"] for s in in_unit if s["side"] != "pg"]
            tdp = lambda b: b.get("tdp", b.get("count", 0)) or 0
            pg_sku = sum(b.get("sku_count", 0) for b, side in brands if side == "pg")
            all_sku = sum(b.get("sku_count", 0) for b, _ in brands)
            pg_tdp = sum(tdp(b) for b, side in brands if side == "pg")
            all_tdp = sum(tdp(b) for b, _ in brands)
            pgm = median(pg_vals) if pg_vals else None
            cpm = median(comp_vals) if comp_vals else None
            cat_rows.append([
                cat, retailer_name(rid), pg_sku, all_sku, pg_sku / all_sku if all_sku else None,
                None if national else pg_tdp, None if national else all_tdp,
                None if national or not all_tdp else pg_tdp / all_tdp,
                pgm, cpm, UNIT_LABEL.get(unit, unit or ""), (pgm / cpm - 1) if pgm and cpm else None,
            ])
            for b, side in brands:
                vals = [s["unit_price"] for s in in_unit if _brand_match(s, b["brand"], side)]
                brand_rows.append([
                    cat, retailer_name(rid), "P&G" if side == "pg" else "Competitor", b["brand"],
                    b.get("sku_count", 0), None if national else b.get("store_count"),
                    None if national else tdp(b),
                    None if national or not n_stores else (b.get("store_count") or 0) / n_stores,
                    median(vals) if vals else None, UNIT_LABEL.get(unit, unit or "") if vals else "",
                    (b.get("promo_pct") or 0) / 100,
                ])
    _add_table(wb, "Categories", ["Category", "Retailer", "P&G products", "Tracked products", "P&G share of products",
                                  "P&G TDP", "Tracked TDP", "P&G share of TDP", "P&G typical price per unit",
                                  "Competitor typical price per unit", "Per", "P&G vs competitors"], cat_rows,
               {"P&G share of products": PCT, "P&G TDP": INT, "Tracked TDP": INT, "P&G share of TDP": PCT,
                "P&G typical price per unit": EUR3, "Competitor typical price per unit": EUR3, "P&G vs competitors": PCT},
               note="Tracked = P&G + the named competitors for that category. Price per unit: per wash for Laundry/Fabric Enhancers, otherwise per litre, kg or piece (the category's most common unit).")
    _add_table(wb, "Brands", ["Category", "Retailer", "Side", "Brand", "Products", "Stores listing", "TDP",
                              "% of stores", "Typical price per unit", "Per", "% of listings on promo"], brand_rows,
               {"TDP": INT, "% of stores": PCT, "Typical price per unit": EUR3, "% of listings on promo": PCT})

    # ---------- Promotions ----------
    rows = []
    for r in retailers:
        for p in (r.get("insights") or {}).get("promos", []):
            rows.append([retailer_name(r["retailer"]), "P&G" if p["side"] == "pg" else "Competitor", p["product"],
                         p.get("barcode") if _real_ean(p.get("barcode")) else "", p["brand"], p.get("fighting_category") or "",
                         p.get("regular_price"),
                         "estimated from other stores" if p.get("regular_price_source") == "other_stores" else "",
                         p["promo_price"], (p["depth_pct"] / 100) if p.get("depth_pct") is not None else None,
                         p["store_count"], "one national list" if _is_national(r) else r.get("stores_count")])
    _add_table(wb, "Promotions", ["Retailer", "Side", "Product", "EAN", "Brand", "Category", "Regular price",
                                  "Regular price note", "Promo price", "Discount", "Stores running it", "Stores reporting"],
               rows, {"Regular price": EUR, "Promo price": EUR, "Discount": PCT}, {"Product": 42, "EAN": 15})

    # ---------- Missing products ----------
    rows = []
    for rid in rids:
        for f in flags[rid]:
            rows.append([retailer_name(rid), f.get("product"), f.get("barcode") if _real_ean(f.get("barcode")) else "",
                         f.get("brand"), f.get("fighting_category") or f.get("category") or "", f.get("store_name"),
                         f.get("city"), f.get("last_seen_date"), f.get("last_seen_price"), f.get("days_missing"),
                         "Yes" if f.get("was_on_promo") else ""])
    _add_table(wb, "Missing products", ["Retailer", "Product", "EAN", "Brand", "Category", "Store", "City",
                                        "Last listed", "Last price", "Days missing", "On promo when last listed"],
               rows, {"Last price": EUR}, {"Product": 42, "EAN": 15, "Store": 26},
               note="Listed by the store before, absent 3+ days in a row on days that store published. A signal to check with the store, not a confirmed out-of-stock.")

    # ---------- Store portfolio ----------
    portfolio_rows, store_sheets = [], []
    for r in retailers:
        rid = r["retailer"]
        products = _load(DOCS_DIR / "products" / f"{rid}.json", [])
        if not _is_national(r):  # a national list has no store-by-store view
            store_sheets.append((rid, products))
        if _is_national(r) or not products:
            continue
        stores: dict[str, dict] = {}
        chain = defaultdict(set)
        for x in products:
            s = stores.setdefault(x["store_id"], {"name": x["store_name"], "city": x.get("city") or "",
                                                  "addr": x.get("street_address") or "", "here": set(), "promo": 0})
            k = x.get("barcode") or x["product"]
            chain[k].add(x["store_id"])
            s["here"].add(k)
            s["promo"] += x.get("promo_price") is not None
        core = {k for k, ss in chain.items() if len(ss) >= len(stores) / 2}
        flagged = defaultdict(int)
        for f in flags[rid]:
            flagged[f["store_id"]] += 1
        for sid, s in sorted(stores.items(), key=lambda kv: (kv[1]["city"], kv[1]["addr"])):
            present = len(core & s["here"])
            portfolio_rows.append([retailer_name(rid), s["name"], s["city"], s["addr"], len(s["here"]), len(core),
                                   present, present / len(core) if core else None, len(core) - present,
                                   s["promo"], flagged.get(sid, 0)])
    _add_table(wb, "Store portfolio", ["Retailer", "Store", "City", "Address", "P&G products listed", "Core products (chain)",
                                       "Core present", "Core present %", "Core missing", "P&G listings on promo",
                                       "P&G products flagged missing"], portfolio_rows, {"Core present %": PCT},
               {"Store": 26, "Address": 30},
               note="Core product = listed in at least half of that retailer's stores today (dashboard rule, not an official must-stock list).")

    # ---------- per retailer ----------
    for rid in rids:
        rows = [[s["ean"] if _real_ean(s["ean"]) else "", s["product"], s["brand"], "P&G" if s["side"] == "pg" else "Competitor",
                 s.get("cat") or "", s["stores"], s.get("shelf"),
                 s.get("unit_price") if s.get("unit") != "pak" else None,
                 UNIT_LABEL.get(s.get("unit"), s.get("unit") or "") if s.get("unit") != "pak" else "",
                 s.get("promo"), s.get("promo_stores")]
                for s in sorted(skus[rid], key=lambda s: (s["side"] != "pg", s["brand"] or "", s["product"]))]
        _add_table(wb, f"{retailer_name(rid)} SKUs", ["EAN", "Product", "Brand", "Side", "Category", "Stores",
                                                      "Regular price", "Price per unit", "Per", "Promo price", "Stores on promo"],
                   rows, {"Regular price": EUR, "Price per unit": EUR3, "Promo price": EUR}, {"Product": 42, "EAN": 15},
                   note="Every P&G and tracked competitor product at this retailer today. Regular price = typical across its stores.")
    for rid, products in store_sheets:
        rows = [[x.get("store_name"), x.get("city"), x.get("street_address"), x.get("product"),
                 x.get("barcode") if _real_ean(x.get("barcode")) else "", x.get("pg_brand") or x.get("brand"),
                 x.get("fighting_category") or x.get("category"),
                 x.get("price") if x.get("promo_price") is None else x.get("regular_price"),
                 x.get("promo_price"), x.get("unit_price")] for x in products]
        _add_table(wb, f"{retailer_name(rid)} by store", ["Store", "City", "Address", "Product", "EAN", "Brand", "Category",
                                                          "Regular price", "Promo price", "Unit price (as published)"],
                   rows, {"Regular price": EUR, "Promo price": EUR, "Unit price (as published)": EUR},
                   {"Store": 24, "Address": 28, "Product": 42, "EAN": 15},
                   note="Every P&G listing, store by store. Regular price is blank on promo rows where the retailer doesn't publish it.")

    DOCS_DIR.mkdir(parents=True, exist_ok=True)
    wb.save(REPORT_PATH)
    print(f"Wrote {REPORT_PATH}")


if __name__ == "__main__":
    build_report()
