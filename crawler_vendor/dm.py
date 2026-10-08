import datetime
import json
import logging
import re
import warnings
from io import BytesIO
from tempfile import TemporaryFile
from typing import Any

import openpyxl

from .base import BaseCrawler
from .models import Product, Store

logger = logging.getLogger(__name__)


class DmCrawler(BaseCrawler):
    """
    Crawler for dm (drogerie markt) Croatia.

    Adapted from the cijene-api project. dm publishes ONE national price list
    per day (an Excel file linked from its "važeći cjenik" page), valid for
    all its stores and the online shop, rather than one file per store. It is
    therefore returned as a single store with id "all": prices, promotions
    and the cross-retailer comparison work as for any retailer, but
    store-level views (distribution across stores, missing products per
    store, store check) have nothing to compare for dm.
    """

    CHAIN = "dm"
    BASE_URL = "https://www.dm.hr"
    CONTENT_BASE_URL = "https://content.services.dmtech.com/rootpage-dm-shop-hr-hr"
    INDEX_URL = f"{CONTENT_BASE_URL}/novo/promocije/nove-oznake-cijena-i-vazeci-cjenik-u-dm-u-2906632?mrclx=false"

    STORE_ID = "all"
    STORE_NAME = "dm (national price list)"

    # parse_excel reads columns by their (normalized) header text instead of
    # the shared map-driven CSV parser, so these stay empty.
    PRICE_MAP = {}
    FIELD_MAP = {}
    REQUIRED_COLUMNS = []

    # Normalized header text (lower case, no diacritics) -> Product field.
    # Matched by prefix, because dm's headers are long phrases that have
    # changed wording before (e.g. the special-price column's explanation).
    COLUMN_PREFIXES = {
        "naziv": "product",
        "sifra": "product_id",
        "marka": "brand",
        "barkod": "barcode",
        "kategorija": "category",
        "neto kolicina": "quantity",
        "jedinica mjere": "unit",
        "cijena za jedinicu mjere": "unit_price",
        "mpc za vrijeme posebnog oblika prodaje": "special_price",
        "najniza cijena u posljednjih 30 dana": "best_price_30",
        "sidrena cijena": "anchor_price",
        "mpc": "price",
    }
    PRICE_FIELDS = {"price", "unit_price", "special_price", "best_price_30", "anchor_price"}

    @staticmethod
    def parse_date_from_title(title: str) -> datetime.date:
        m = re.search(r"(\d{1,2})\.(\d{1,2})\.(\d{4})", title)
        if not m:
            raise ValueError(f"Could not extract date from title: {title}")
        day, month, year = map(int, m.groups())
        return datetime.date(year, month, day)

    def find_excel_url(self, json_content: str, target_date: datetime.date) -> str:
        data = json.loads(json_content)
        entries = [i.get("data", {}) for i in data.get("mainData", []) if i.get("type") == "CMDownload"]
        if not entries:
            raise ValueError("No Excel links found on dm's price list page")
        for entry in entries:
            headline, link = entry.get("headline", ""), entry.get("linkTarget", "")
            if not headline or not link:
                continue
            try:
                if self.parse_date_from_title(headline) == target_date:
                    return link if link.startswith(("http://", "https://")) else f"{self.CONTENT_BASE_URL}{link}"
            except ValueError:
                continue
        raise ValueError(f"No dm price list published for {target_date:%d.%m.%Y}")

    def _norm_header(self, value: Any) -> str:
        return " ".join(self.strip_diacritics(str(value or "").lower()).split())

    def detect_columns(self, worksheet: Any) -> tuple[int, dict[int, str]]:
        """
        Find the header row and map column index -> Product field.

        dm's header has a merged "naziv + šifra" cell spanning two data
        columns (name, then product code).
        """
        for row_idx, row in enumerate(worksheet.iter_rows(max_row=30), start=1):
            names = [self._norm_header(c.value) for c in row]
            if "naziv + sifra" not in names:
                continue
            cols: dict[int, str] = {}
            for i, name in enumerate(names):
                if name == "naziv + sifra":
                    cols[i], cols[i + 1] = "product", "product_id"
                    continue
                if not name or i in cols:
                    continue
                # longest matching prefix wins ("mpc za vrijeme…" before "mpc")
                for prefix in sorted(self.COLUMN_PREFIXES, key=len, reverse=True):
                    if name.startswith(prefix):
                        field = self.COLUMN_PREFIXES[prefix]
                        if field not in cols.values():
                            cols[i] = field
                        break
            missing = {"product", "product_id", "price", "barcode"} - set(cols.values())
            if missing:
                raise ValueError(f"dm price list is missing columns: {sorted(missing)} (headers: {names})")
            return row_idx, cols
        raise ValueError("Could not find the header row in dm's price list — the file format may have changed")

    def parse_excel(self, excel_data: bytes) -> list[Product]:
        with warnings.catch_warnings():
            warnings.simplefilter("ignore", UserWarning)
            workbook = openpyxl.load_workbook(BytesIO(excel_data), data_only=True, read_only=True)
        worksheet = workbook.active
        header_row, cols = self.detect_columns(worksheet)

        products: list[Product] = []
        for row_idx, row in enumerate(worksheet.iter_rows(min_row=header_row + 1, values_only=True), start=header_row + 1):
            data: dict[str, Any] = {f: "" for f in ("product", "product_id", "brand", "barcode", "category", "quantity", "unit")}
            for i, field in cols.items():
                value = row[i] if i < len(row) else None
                if field in self.PRICE_FIELDS:
                    data[field] = self.parse_price(str(value) if value is not None else None, False)
                else:
                    data[field] = str(value if value is not None else "").strip()
            if not data["product_id"] or not data["product"]:
                continue
            for f in self.PRICE_FIELDS:
                data.setdefault(f, None)
            try:
                products.append(Product(**self.fix_product_data(data)))
            except Exception as e:
                logger.warning(f"dm: skipping row {row_idx}: {e}")
        logger.info(f"dm: parsed {len(products)} products")
        return products

    def get_all_products(self, date: datetime.date) -> list[Store]:
        excel_url = self.find_excel_url(self.fetch_text(self.INDEX_URL), date)
        logger.info(f"dm: price list for {date}: {excel_url}")
        with TemporaryFile(mode="w+b") as tmp:
            self.fetch_binary(excel_url, tmp)
            tmp.seek(0)
            products = self.parse_excel(tmp.read())
        if not products:
            return []
        return [
            Store(
                chain=self.CHAIN,
                store_type="national",
                store_id=self.STORE_ID,
                name=self.STORE_NAME,
                street_address="",
                zipcode="",
                city="",
                items=products,
            )
        ]
